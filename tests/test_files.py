"""Folder tidy and the drawing register (filesystem / openpyxl, no SolidWorks)."""
import os
import shutil
import subprocess

import pytest

from swdraw.numbering import Scheme
from swdraw.register import build_register
from swdraw.tidy import tidy


def write(p, data):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_bytes(data)
    return p


def test_tidy_dry_run_changes_nothing(tmp_path):
    w = tmp_path / "Job" / "Working"
    write(w / "LIVE.SLDPRT", b"live")
    write(w / "old.SLDPRT", b"old")
    rep = tidy(str(w), keep=["live.*"])
    assert rep["dry_run"] and any("old.SLDPRT" in m for m in rep["moves"])
    assert (w / "old.SLDPRT").exists()


def test_tidy_moves_merges_and_suffixes(tmp_path):
    job = tmp_path / "Job"
    w, sup = job / "Working", job / "_Superseded"
    write(w / "LIVE.SLDDRW", b"live")
    write(w / "same.SLDPRT", b"same")
    write(sup / "same.SLDPRT", b"same")
    write(w / "diff.SLDPRT", b"new")
    write(sup / "diff.SLDPRT", b"older")
    write(w / "~$Register.xlsx", b"lock")
    (w / "Reference").mkdir()
    write(w / "Reference" / "r.pdf", b"r")
    (w / "Scratch").mkdir()
    rep = tidy(str(w), keep=["LIVE.*"], move_folders={"Reference": "<parent>"}, dry_run=False)
    assert sorted(os.listdir(w)) == ["LIVE.SLDDRW", "Scratch", "~$Register.xlsx"]
    assert (sup / "same.SLDPRT").read_bytes() == b"same"
    assert (sup / "diff.SLDPRT").read_bytes() == b"older"                   # never clobbered
    assert any(n.startswith("diff (") for n in os.listdir(sup))
    assert (job / "Reference" / "r.pdf").exists()
    assert any("merged" in m for m in rep["moves"])
    assert "folder Scratch" in rep["skipped"] and any("lock" in s for s in rep["skipped"])


def test_tidy_missing_folder(tmp_path):
    assert tidy(str(tmp_path / "nope"), keep=[])["errors"]


DATA = {
    "title": "ACME1 drawing register",
    "drawings": [
        {"number": "ACME1-MEC-2100", "title": "BUSH ASSEMBLY", "type": "Assembly"},
        {"number": "ACME1-MEC-2101", "title": "FLANGED BUSH", "type": "Part", "parent": "ACME1-MEC-2100"},
    ],
    "revisions": [
        {"number": "ACME1-MEC-2101", "rev": "A", "date": "2026-09-01", "description": "FIRST ISSUE", "status": "PROTOTYPE"},
        {"number": "ACME1-MEC-2100", "rev": "A", "date": "01/09/2026", "description": "FIRST ISSUE", "status": "DRAFT"},
        {"number": "ACME1-MEC-2101", "rev": "B", "date": "2026-10-18", "description": "BORE TOL", "status": "RELEASED"},
    ],
}


def test_register_structure(tmp_path):
    from openpyxl import load_workbook
    path = build_register(str(tmp_path / "reg.xlsx"), DATA, Scheme())
    wb = load_workbook(path)
    assert wb.sheetnames == ["Register", "Revision History", "Numbering"]
    reg, hist = wb["Register"], wb["Revision History"]
    assert reg["A6"].value == "ACME1-MEC-2101"
    assert reg["E6"].value.startswith("=IFERROR(LOOKUP(2,1/(")
    assert hist["C2"].number_format == "DD/MM/YYYY"
    assert hist["A2"].fill.fgColor.rgb.endswith("FFF2CC")
    assert len(hist.data_validations.dataValidation) == 2
    num = wb["Numbering"]
    kinds = [c.value for c in num["B"] if c.value]
    assert any("assembly" in str(k) for k in kinds)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_register_formulas_calculate(tmp_path):
    """Recalculate with LibreOffice and check the latest revision wins."""
    from openpyxl import load_workbook
    src = build_register(str(tmp_path / "reg.xlsx"), DATA, Scheme())
    out = tmp_path / "calc"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), src],
                   check=True, capture_output=True, timeout=120)
    wb = load_workbook(out / "reg.xlsx", data_only=True)
    reg = wb["Register"]
    assert reg["E5"].value == "A" and reg["G5"].value == "DRAFT"
    assert reg["E6"].value == "B" and reg["G6"].value == "RELEASED"
    d = reg["F6"].value
    assert (d.year, d.month, d.day) == (2026, 10, 18)
